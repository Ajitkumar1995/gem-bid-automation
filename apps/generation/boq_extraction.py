"""Conservative extraction of tabular BoQ data from bid documents."""
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import re
from pathlib import Path
import tempfile
from zipfile import ZipFile, is_zipfile

from django.db import transaction


MAX_PDF_PAGES = 100
MAX_WORKSHEETS = 50
MAX_WORKSHEET_ROWS = 10000
MAX_WORKSHEET_COLUMNS = 100

_DESCRIPTION_HEADER = re.compile(
    r"\b(item|description|product|service|material|work|particular|goods?|"
    r"specification|scope|name)\b",
    re.IGNORECASE,
)
_ITEM_IDENTIFIER_HEADER = re.compile(r"\b(item\s*(no|number|code|id)|serial|sl\s*(no|number))\b", re.IGNORECASE)
_QUANTITY_HEADER = re.compile(r"\b(qty|quantity|nos|number of units|required quantity)\b", re.IGNORECASE)
_UNIT_HEADER = re.compile(r"\b(unit|uom|unit of measure|measurement)\b", re.IGNORECASE)
_RESPONSE_HEADER = re.compile(r"\b(bidder response|compliance|offered|remarks?|response)\b", re.IGNORECASE)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _unique_headers(values):
    headers = []
    counts = {}
    for index, value in enumerate(values):
        header = _cell_text(value) or f"Column {index + 1}"
        counts[header] = counts.get(header, 0) + 1
        if counts[header] > 1:
            header = f"{header} ({counts[header]})"
        headers.append(header[:120])
    return headers


def _header_mapping(rows):
    best = None
    for header_index, row in enumerate(rows[:20]):
        headers = _unique_headers(row[:MAX_WORKSHEET_COLUMNS])
        description_indexes = [
            i for i, header in enumerate(headers)
            if _DESCRIPTION_HEADER.search(header) and not _ITEM_IDENTIFIER_HEADER.search(header)
        ]
        quantity_indexes = [i for i, header in enumerate(headers) if _QUANTITY_HEADER.search(header)]
        unit_indexes = [i for i, header in enumerate(headers) if _UNIT_HEADER.search(header)]
        response_indexes = [i for i, header in enumerate(headers) if _RESPONSE_HEADER.search(header)]
        score = bool(description_indexes) + bool(quantity_indexes) + bool(unit_indexes) + bool(response_indexes)
        if description_indexes and quantity_indexes and score >= 2 and (best is None or score > best[0]):
            description_indexes.sort(key=lambda i: (
                not bool(re.search(r"description|particular|product|service|material|work|name|item", headers[i], re.IGNORECASE)),
                bool(re.search(r"specification", headers[i], re.IGNORECASE)),
            ))
            best = (score, header_index, headers, description_indexes[0], quantity_indexes[0],
                    unit_indexes[0] if unit_indexes else None,
                    response_indexes[0] if response_indexes else None)
    return best


def _normalized_table(rows, source_label):
    clean_rows = [
        [_cell_text(cell) for cell in row[:MAX_WORKSHEET_COLUMNS]]
        for row in rows
        if any(_cell_text(cell) for cell in row[:MAX_WORKSHEET_COLUMNS])
    ]
    if not clean_rows:
        return None
    mapping = _header_mapping(clean_rows)
    if mapping is None:
        return None
    _score, header_index, headers, description_index, quantity_index, unit_index, response_index = mapping
    items = []
    for row_number, row in enumerate(clean_rows[header_index + 1:], start=header_index + 1):
        values = (row + [""] * len(headers))[:len(headers)]
        if [value.casefold() for value in values] == [header.casefold() for header in headers]:
            continue
        description = values[description_index].strip()
        quantity = values[quantity_index].strip()
        if not description or not quantity:
            continue
        raw = {header: value for header, value in zip(headers, values) if value}
        items.append({
            "description": description[:500],
            "quantity": quantity[:50],
            "unit": values[unit_index][:60] if unit_index is not None else "",
            "bidder_response": (
                values[response_index][:120].strip() or "Complied"
                if response_index is not None
                else "Complied"
            ),
            "source_data": raw,
            "source_row": row_number,
        })
    if not items:
        return None
    return {"source_label": source_label[:255], "headers": headers, "items": items}


def _read_spreadsheet(content, filename):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    tables = []
    try:
        for worksheet in workbook.worksheets[:MAX_WORKSHEETS]:
            rows = list(worksheet.iter_rows(
                min_row=1,
                max_row=MAX_WORKSHEET_ROWS,
                max_col=MAX_WORKSHEET_COLUMNS,
                values_only=True,
            ))
            table = _normalized_table(rows, f"{filename} · {worksheet.title}")
            if table:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def _read_legacy_excel(content, filename):
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    tables = []
    try:
        for worksheet in workbook.sheets()[:MAX_WORKSHEETS]:
            rows = [
                worksheet.row_values(row_index, start_colx=0, end_colx=min(worksheet.ncols, MAX_WORKSHEET_COLUMNS))
                for row_index in range(min(worksheet.nrows, MAX_WORKSHEET_ROWS))
            ]
            table = _normalized_table(rows, f"{filename} · {worksheet.name}")
            if table:
                tables.append(table)
    finally:
        workbook.release_resources()
    return tables


def _read_csv(content, filename):
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(StringIO(text), dialect))
    table = _normalized_table(rows, filename)
    return [table] if table else []


def _read_pdf(content, filename):
    import camelot
    import pymupdf

    with pymupdf.open(stream=content, filetype="pdf") as pdf:
        page_count = min(len(pdf), MAX_PDF_PAGES)
    if not page_count:
        return []
    with tempfile.NamedTemporaryFile(suffix=".pdf") as temp_pdf:
        temp_pdf.write(content)
        temp_pdf.flush()
        extracted = camelot.read_pdf(
            temp_pdf.name,
            pages=f"1-{page_count}",
            flavor="stream",
        )
    tables = []
    for index, candidate in enumerate(extracted, start=1):
        table = _normalized_table(candidate.df.values.tolist(), f"{filename} · table {index}")
        if table:
            tables.append(table)
    return tables


def extract_boq_tables(content, filename, content_type=""):
    """Return only tables with both recognizable item-description and quantity headers."""
    suffix = Path(filename).suffix.lower()
    normalized_type = (content_type or "").split(";", 1)[0].lower()
    is_xlsx_content = False
    is_zip_content = content.startswith(b"PK") and is_zipfile(BytesIO(content))
    if is_zip_content:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
        is_xlsx_content = {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names)
    if is_xlsx_content or suffix in {".xlsx", ".xlsm"} or normalized_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroenabled.12",
    }:
        return _read_spreadsheet(content, filename)
    if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") or suffix == ".xls":
        return _read_legacy_excel(content, filename)
    if content.startswith(b"%PDF-") or normalized_type == "application/pdf":
        return _read_pdf(content, filename)
    if suffix == ".csv" or normalized_type in {"text/csv", "application/csv"}:
        return _read_csv(content, filename)
    if suffix == ".pdf":
        return _read_pdf(content, filename)
    if not is_zip_content and b"\x00" not in content:
        return _read_csv(content, filename)
    return []


@transaction.atomic
def import_boq_tables(preparation, tables):
    """Add extracted line items once without replacing bidder-entered rows."""
    from .models import BidBOQItem

    created = 0
    for table in tables:
        for item in table["items"]:
            source_key = sha256(
                repr((
                    table["source_label"],
                    item["source_row"],
                    sorted(item["source_data"].items()),
                )).encode("utf-8")
            ).hexdigest()
            _row, was_created = BidBOQItem.objects.get_or_create(
                preparation=preparation,
                source_key=source_key,
                defaults={
                    "sort_order": preparation.boq_items.count(),
                    "description": item["description"],
                    "quantity": _quantity_decimal(item["quantity"]),
                    "unit": item["unit"],
                    "bidder_response": item["bidder_response"] or "Complied",
                    "source_data": item["source_data"],
                    "source_label": table["source_label"],
                },
            )
            created += int(was_created)
    return created


def _quantity_decimal(value):
    from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

    normalized = value.replace(",", "").strip()
    try:
        quantity = Decimal(normalized).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None
    if not quantity.is_finite() or abs(quantity) >= Decimal("100000000000"):
        return None
    return quantity


def extracted_boq_tables(preparation):
    """Build display-ready full-source tables from imported rows."""
    groups = {}
    for item in preparation.boq_items.order_by("sort_order", "id"):
        if not item.source_data:
            continue
        group = groups.setdefault(item.source_label, {"source": item.source_label, "headers": [], "rows": []})
        for header in item.source_data:
            if header not in group["headers"]:
                group["headers"].append(header)
        group["rows"].append(item)
    result = []
    for group in groups.values():
        headers = [*group["headers"], "Bidder Response"]
        group["rows"] = [
            {
                "cells": [
                    *[item.source_data.get(header, "") for header in group["headers"]],
                    item.bidder_response or "Complied",
                ],
                "description": item.description,
            }
            for item in group["rows"]
        ]
        group["headers"] = headers
        result.append(group)
    return result


def import_cached_linked_boq_tables(preparation, max_files=10):
    """Recover BoQ rows from downloaded CSV and spreadsheet links."""
    from apps.documents.models import LinkedDocument

    supported_suffixes = {".csv", ".xls", ".xlsx", ".xlsm"}
    candidates = []
    for document in preparation.bid.linked_documents.filter(
        status=LinkedDocument.Status.DOWNLOADED,
    ).exclude(file=""):
        filename = document.file_name or document.filename
        if Path(filename).suffix.lower() not in supported_suffixes:
            continue
        priority_text = " ".join((filename, document.label, document.url)).casefold()
        candidates.append((
            not bool(re.search(r"\bboq\b|bill.?of.?quantit|line.?items?", priority_text)),
            document.created_at,
            document.pk,
            document,
            filename,
        ))

    imported = 0
    for _not_boq_named, _created_at, _document_id, document, filename in sorted(candidates)[:max_files]:
        with document.file.open("rb") as source:
            content = source.read()
        tables = extract_boq_tables(content, filename, document.content_type)
        imported += import_boq_tables(preparation, tables)
    return imported
